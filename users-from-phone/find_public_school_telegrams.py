#!/usr/bin/env python3
"""Find public Telegram links on the supplied schools' public source pages."""

from __future__ import annotations

import argparse
import asyncio
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from telethon import TelegramClient, errors, types


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_API_ID = 34825825
DEFAULT_API_HASH = "60176f7ad0bcd77e63d4a64ca8d50a38"
URL_RE = re.compile(r"https?://[^\s;,<>\]})]+", re.IGNORECASE)
TG_HOSTS = {"t.me", "www.t.me", "telegram.me", "www.telegram.me"}


@dataclass
class PageResult:
    url: str
    status: str
    telegrams: list[str]


def unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(values))


def source_urls(value: object) -> list[str]:
    return unique(URL_RE.findall(str(value or "")))


def canonical_telegram_url(url: str) -> str | None:
    parsed = urlparse(url)
    if parsed.netloc.lower() in TG_HOSTS:
        path = parsed.path.strip("/")
        if path and not path.startswith(("joinchat/", "addstickers/", "share/")):
            username = path.split("/", 1)[0]
            if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{4,31}", username):
                return f"https://t.me/{username.lower()}"
    if parsed.scheme == "tg" and parsed.netloc == "resolve":
        username = parse_qs(parsed.query).get("domain", [""])[0]
        if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{4,31}", username):
            return f"https://t.me/{username.lower()}"
    return None


def fetch_page(url: str) -> PageResult:
    if canonical_telegram_url(url):
        return PageResult(url, "direct_telegram_link", [canonical_telegram_url(url)])
    try:
        response = requests.get(
            url,
            headers={"User-Agent": "Mozilla/5.0 (compatible; public-contact-audit/1.0)"},
            timeout=(5, 15),
            allow_redirects=True,
        )
        response.raise_for_status()
        if "html" not in response.headers.get("content-type", "").lower():
            return PageResult(url, "not_html", [])
        soup = BeautifulSoup(response.text, "html.parser")
        links = [anchor.get("href", "") for anchor in soup.find_all("a", href=True)]
        links.extend(URL_RE.findall(response.text))
        telegrams = unique([candidate for link in links if (candidate := canonical_telegram_url(link))])
        return PageResult(url, "ok", telegrams)
    except requests.Timeout:
        return PageResult(url, "timeout", [])
    except requests.RequestException as exc:
        return PageResult(url, f"http_error_{getattr(exc.response, 'status_code', 'unknown')}", [])


def account_config(account: str) -> tuple[int, str, Path]:
    prefix = re.sub(r"[^A-Za-z0-9]+", "_", account).strip("_").upper()
    session = os.getenv(f"TG_{prefix}_SESSION") or f"sessions/{account}.session"
    session_path = Path(session)
    if not session_path.is_absolute():
        session_path = BASE_DIR / session_path
    return DEFAULT_API_ID, DEFAULT_API_HASH, session_path


async def check_telegram(client: TelegramClient, url: str) -> tuple[str, str, str]:
    username = url.rsplit("/", 1)[-1]
    try:
        entity = await client.get_entity(username)
    except (errors.UsernameInvalidError, errors.UsernameNotOccupiedError, ValueError):
        return "not_found", "", ""
    except errors.FloodWaitError as exc:
        if exc.seconds > 300:
            raise SystemExit(f"Telegram limited public-link verification for {exc.seconds}s; rerun later or use --skip-telegram-verify.") from exc
        print(f"Telegram asked to wait {exc.seconds}s while checking public links; waiting.", flush=True)
        await asyncio.sleep(exc.seconds + 1)
        return await check_telegram(client, url)
    if isinstance(entity, types.User):
        return "found", "user", entity.first_name or entity.username or ""
    if isinstance(entity, (types.Channel, types.Chat)):
        return "found", "channel_or_group", getattr(entity, "title", "") or ""
    return "found", type(entity).__name__, ""


async def main() -> None:
    parser = argparse.ArgumentParser(description="Find and verify public Telegram links on schools' public sites.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, default=BASE_DIR / "repetam_online_school_leads_public_telegram_checked.xlsx")
    parser.add_argument("--account", default="main18")
    parser.add_argument("--workers", type=int, default=10)
    parser.add_argument("--skip-telegram-verify", action="store_true", help="save public t.me links without resolving them through Telegram")
    args = parser.parse_args()
    if not args.input.is_file():
        raise SystemExit(f"Input file not found: {args.input}")

    source_book = load_workbook(args.input, read_only=True, data_only=False)
    source_sheet = source_book.active
    headers = [cell.value for cell in next(source_sheet.iter_rows(min_row=1, max_row=1))]
    try:
        source_index = headers.index("Источник")
    except ValueError as exc:
        raise SystemExit("Input file has no 'Источник' column.") from exc
    rows = [list(row) for row in source_sheet.iter_rows(min_row=2, values_only=True)]
    all_urls = unique([url for row in rows for url in source_urls(row[source_index])])
    print(f"Checking {len(all_urls)} public source URLs.", flush=True)
    page_results: dict[str, PageResult] = {}
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        futures = {pool.submit(fetch_page, url): url for url in all_urls}
        for number, future in enumerate(as_completed(futures), start=1):
            result = future.result()
            page_results[result.url] = result
            if number % 50 == 0 or number == len(futures):
                print(f"Read {number}/{len(futures)} public pages.", flush=True)

    all_telegrams = unique([telegram for result in page_results.values() for telegram in result.telegrams])
    if args.skip_telegram_verify:
        print(f"Saving {len(all_telegrams)} public Telegram links without direct Telegram verification.", flush=True)
        telegram_results = {telegram: ("not_checked_tg_rate_limit", "", "") for telegram in all_telegrams}
    else:
        print(f"Verifying {len(all_telegrams)} public Telegram links.", flush=True)
        api_id, api_hash, session_path = account_config(args.account)
        client = TelegramClient(str(session_path.with_suffix("")), api_id, api_hash)
        await client.connect()
        if not await client.is_user_authorized():
            await client.disconnect()
            raise SystemExit("Saved Telegram session is unavailable.")
        try:
            telegram_results = {}
            for number, telegram in enumerate(all_telegrams, start=1):
                telegram_results[telegram] = await check_telegram(client, telegram)
                if number % 50 == 0 or number == len(all_telegrams):
                    print(f"Verified {number}/{len(all_telegrams)} Telegram links.", flush=True)
                await asyncio.sleep(0.25)
        finally:
            await client.disconnect()

    output_book = Workbook()
    output_sheet = output_book.active
    output_sheet.title = "Публичные Telegram"
    output_headers = headers + [
        "Публичные Telegram-ссылки", "Страница-источник Telegram", "Статус Telegram-ссылки",
        "Тип Telegram", "Название Telegram", "Статус проверки страниц",
    ]
    output_sheet.append(output_headers)
    for row in rows:
        urls = source_urls(row[source_index])
        telegram_to_sources: dict[str, list[str]] = {}
        page_statuses = []
        for url in urls:
            result = page_results[url]
            page_statuses.append(f"{url}: {result.status}")
            for telegram in result.telegrams:
                telegram_to_sources.setdefault(telegram, []).append(url)
        telegrams = list(telegram_to_sources)
        statuses = [f"{telegram}: {telegram_results[telegram][0]}" for telegram in telegrams]
        kinds = [f"{telegram}: {telegram_results[telegram][1]}" for telegram in telegrams]
        titles = [f"{telegram}: {telegram_results[telegram][2]}" for telegram in telegrams]
        output_sheet.append(row + [
            "; ".join(telegrams), "; ".join(f"{telegram}: {', '.join(telegram_to_sources[telegram])}" for telegram in telegrams),
            "; ".join(statuses), "; ".join(kinds), "; ".join(titles), "; ".join(page_statuses),
        ])
    output_sheet.freeze_panes = "A2"
    for column in output_sheet.columns:
        output_sheet.column_dimensions[column[0].column_letter].width = min(max(12, max(len(str(cell.value or "")) for cell in column) + 2), 55)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    output_book.save(args.output)
    print(f"Saved {args.output}", flush=True)


if __name__ == "__main__":
    asyncio.run(main())
